from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
CHART_FORMULA_RE = re.compile(r"<(?:\w+:)?f>(.*?)</(?:\w+:)?f>")
CELL_RANGE_RE = re.compile(r"^(?:'((?:[^']|'')+)'|([^!]+))!(\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)$")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def cell_values(ws: Any, ref: str) -> list[Any]:
    min_col, min_row, max_col, max_row = range_boundaries(ref.replace("$", ""))
    return [ws.cell(r, c).value for r in range(min_row, max_row + 1) for c in range(min_col, max_col + 1)]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    workbook_path = args.workbook.resolve()
    project_root = next(parent for parent in workbook_path.parents if parent.name == "TFM")
    normalized = project_root / "04_EVIDENCE" / "Excel_visibilidad_26062026" / "normalized"

    formula_wb = load_workbook(workbook_path, data_only=False, keep_links=True)
    value_wb = load_workbook(workbook_path, data_only=True, keep_links=True)

    pending: list[str] = []
    formula_errors: list[str] = []
    stored_errors: list[str] = []
    forbidden_hits: list[dict[str, str]] = []
    allowed_exclusions = {"09_Resultados_Custom!A2", "README!B17", "00_Guia!A2"}
    forbidden_patterns = (
        "benchmark",
        "cpu media",
        "working set",
        "private bytes",
        "consumo por proceso",
        "b0-b8",
        "b0–b8",
    )

    for ws in formula_wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str):
                    if value.strip().casefold() == "pendiente":
                        pending.append(f"{ws.title}!{cell.coordinate}")
                    upper = value.upper()
                    if value.startswith("=") and any(token in upper for token in FORMULA_ERRORS):
                        formula_errors.append(f"{ws.title}!{cell.coordinate}:{value}")
                    location = f"{ws.title}!{cell.coordinate}"
                    lower = value.casefold()
                    if location not in allowed_exclusions and any(pattern in lower for pattern in forbidden_patterns):
                        if ws.title not in {"TECNICAS_REAL_VR", "ALERTAS_VR"}:
                            forbidden_hits.append({"cell": location, "value": value[:240]})
                stored = value_wb[ws.title][cell.coordinate].value
                if isinstance(stored, str) and any(token in stored.upper() for token in FORMULA_ERRORS):
                    stored_errors.append(f"{ws.title}!{cell.coordinate}:{stored}")

    with (normalized / "alertas_vr.csv").open(encoding="utf-8-sig", newline="") as handle:
        alerts = list(csv.DictReader(handle))
    alert_profiles = Counter(row["profile"] for row in alerts)
    alert_tecs = Counter(row["TEC"] for row in alerts)

    chart_details: list[dict[str, Any]] = []
    chart_empty_sources: list[str] = []
    chart_broken_refs: list[str] = []
    external_parts: list[str] = []
    xml_errors: list[str] = []
    with zipfile.ZipFile(workbook_path) as archive:
        zip_test = archive.testzip()
        names = archive.namelist()
        external_parts = [name for name in names if name.startswith("xl/externalLinks/")]
        chart_parts = sorted(name for name in names if re.fullmatch(r"xl/charts/chart\d+\.xml", name))
        for part in chart_parts:
            text = archive.read(part).decode("utf-8", errors="replace")
            refs = [re.sub(r"\s+", " ", item).strip() for item in CHART_FORMULA_RE.findall(text)]
            empty_refs: list[str] = []
            for ref in refs:
                if "#REF!" in ref.upper():
                    chart_broken_refs.append(f"{part}:{ref}")
                    continue
                match = CELL_RANGE_RE.match(ref)
                if not match:
                    continue
                sheet_name = (match.group(1) or match.group(2)).replace("''", "'")
                range_ref = match.group(3)
                if sheet_name not in value_wb.sheetnames:
                    chart_broken_refs.append(f"{part}:{ref}")
                    continue
                values = cell_values(value_wb[sheet_name], range_ref)
                if not any(value not in (None, "") for value in values):
                    empty_refs.append(ref)
            if empty_refs:
                chart_empty_sources.extend(f"{part}:{ref}" for ref in empty_refs)
            chart_details.append({"part": part, "refs": refs, "empty_refs": empty_refs})

    checks: dict[str, bool] = {}
    checks["zip_openxml_valid"] = zip_test is None and not xml_errors
    checks["sheet_count_30"] = len(formula_wb.sheetnames) == 30
    checks["benchmark_sheet_absent"] = "09_Benchmark_Plan" not in formula_wb.sheetnames
    checks["results_custom_sheet_present"] = "09_Resultados_Custom" in formula_wb.sheetnames
    checks["pending_exact_zero"] = len(pending) == 0
    checks["formula_errors_zero"] = not formula_errors and not stored_errors
    checks["external_links_zero"] = len(external_parts) == 0 and not getattr(formula_wb, "_external_links", [])
    checks["charts_11"] = len(chart_details) == 11
    checks["chart_broken_refs_zero"] = not chart_broken_refs
    checks["chart_empty_sources_zero"] = not chart_empty_sources
    checks["forbidden_benchmark_analysis_zero"] = not forbidden_hits

    results = value_wb["09_Resultados_Custom"]
    checks["profile_counts_19_118_109_133"] = [results[f"F{row}"].value for row in range(5, 9)] == [19, 118, 109, 133]
    checks["custom_total_379"] = results["F9"].value == 379
    checks["custom_detection_9_of_9"] = results["E9"].value == 9
    checks["custom_fp_hits_zero"] = results["H9"].value == 0 and results["G64"].value == 0
    checks["tec_matrix_9_rows"] = [results[f"A{row}"].value for row in range(13, 22)] == [f"TEC-{i:03d}" for i in range(1, 10)]
    checks["tec_alert_counts_match_csv"] = all(results[f"F{12+i}"].value == alert_tecs[f"TEC-{i:03d}"] for i in range(1, 10))
    checks["alertas_vr_rows_379"] = value_wb["ALERTAS_VR"].max_row - 4 == 379 and len(alerts) == 379
    checks["alert_profile_counts_match_csv"] = alert_profiles == Counter({"P1": 19, "P2": 118, "P3": 109, "P4": 133})
    checks["tec_campaign_9_ok"] = results["B63"].value == 9 and results["D63"].value == 9 and results["E63"].value == 0 and results["F63"].value == 0
    checks["fp_campaign_10_ok"] = results["B64"].value == 10 and results["D64"].value == 10 and results["E64"].value == 0 and results["F64"].value == 0

    dashboard = value_wb["01_Dashboard"]
    checks["dashboard_custom_9_of_9"] = dashboard["B5"].value == "9/9"
    checks["dashboard_public_ch_3_of_9"] = dashboard["T4"].value == 3
    checks["dashboard_wazuh_base_3_of_9"] = dashboard["T5"].value == 3
    checks["dashboard_wazuh_custom_4_of_9"] = dashboard["T6"].value == 4
    checks["public_campaigns_5_positive_1_inconclusive"] = dashboard["T45"].value == 5 and dashboard["T47"].value == 1
    dashboard_hyperlinks = [cell for row in dashboard.iter_rows() for cell in row if cell.hyperlink]
    checks["dashboard_internal_hyperlinks_10"] = len(dashboard_hyperlinks) >= 10

    guide = value_wb["00_Guia"]
    checks["guide_points_to_results_custom"] = guide["A13"].value == "09_Resultados_Custom"
    checks["catalog_custom_4_detectors"] = [value_wb["07_Catalogo_Artifacts"][f"A{r}"].value for r in range(5, 9)] == [
        "Custom.TFM.HIDS.P1.Critical.Priority.Event_v1",
        "Custom.TFM.HIDS.P2.High.Forensic.Event_v1",
        "Custom.TFM.HIDS.P3.Medium.Behavioral.Event_v1",
        "Custom.TFM.HIDS.P4.Low.Basic_v2",
    ]
    checks["source_group_rows_35"] = sum(1 for row in range(25, 60) if results[f"A{row}"].value) == 35
    checks["receiver_http_separate"] = value_wb["VISIBILIDAD_SISTEMA"]["A11"].value == "receiver HTTP"
    checks["raw_alert_data_preserved"] = len(alerts) == 379

    artifact_expected = {
        "Custom.TFM.HIDS.P1.Critical.Priority.Event_v1.yaml": "8A1FF2654405AF9E4412D32EF0C2BE08A383924654AA764A04AB264BAA42EDD7",
        "Custom.TFM.HIDS.P2.High.Forensic.Event_v1.yaml": "255C5F7CA39E362FE45106C662C4C35730EFF57A4063A1E9973FD11A6375FAA2",
        "Custom.TFM.HIDS.P3.Medium.Behavioral.Event_v1.yaml": "C591CB849753194F6F87834C1E0F34EDE6CDFF6B31BC63650A1D281ADA1D410A",
        "Custom.TFM.HIDS.P4.Low.Basic_v2.yaml": "1EBB29296A14475E38BCF34CA8DD9CC16572F4A8CA00A1AB0C3608E0B3097BC1",
        "Custom.TFM.HIDS.Router.JSONL.Discord.SOC_v3.yaml": "574E3E8BC03E7C4D655400D5BB5527771E438CA67C413853A0AAAC7A10C5E33F",
    }
    artifact_actual = {
        name: sha256(project_root / "01_ARTIFACTS" / "validated" / name) for name in artifact_expected
    }
    checks["artifact_hashes_match"] = artifact_actual == artifact_expected

    all_pass = all(checks.values())
    payload = {
        "generated_at": datetime.now(timezone.utc).astimezone().isoformat(),
        "workbook": str(workbook_path),
        "bytes": workbook_path.stat().st_size,
        "sha256": sha256(workbook_path),
        "status": "PASS" if all_pass else "FAIL",
        "checks": checks,
        "counts": {
            "sheets": len(formula_wb.sheetnames),
            "charts": len(chart_details),
            "tables": sum(len(ws.tables) for ws in formula_wb.worksheets),
            "pending_exact": len(pending),
            "formula_errors": len(formula_errors) + len(stored_errors),
            "external_links": len(external_parts),
            "chart_broken_refs": len(chart_broken_refs),
            "chart_empty_sources": len(chart_empty_sources),
            "forbidden_benchmark_hits": len(forbidden_hits),
            "hyperlinks_dashboard": len(dashboard_hyperlinks),
            "alerts": len(alerts),
            "source_result_rows": 35,
        },
        "pending_cells": pending,
        "formula_error_cells": formula_errors + stored_errors,
        "forbidden_benchmark_hits": forbidden_hits,
        "chart_broken_refs": chart_broken_refs,
        "chart_empty_sources": chart_empty_sources,
        "artifact_hashes": artifact_actual,
        "chart_details": chart_details,
    }

    text = json.dumps(payload, ensure_ascii=False, indent=2) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text, encoding="utf-8")
    print(json.dumps({"status": payload["status"], "failed": [key for key, value in checks.items() if not value], "counts": payload["counts"]}, ensure_ascii=False))
    formula_wb.close()
    value_wb.close()
    return 0 if all_pass else 1


if __name__ == "__main__":
    raise SystemExit(main())
