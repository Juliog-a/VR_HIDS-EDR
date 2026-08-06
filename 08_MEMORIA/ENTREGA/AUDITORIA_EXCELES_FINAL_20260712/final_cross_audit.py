from __future__ import annotations

import hashlib
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(r"C:\Users\julio\Desktop\TFM")
ENTREGABLE = ROOT / "08_MEMORIA" / "ENTREGABLE"
AUDIT_DIR = ROOT / "08_MEMORIA" / "AUDITORIA_EXCELES_FINAL_20260712"
VIS_MASTER = ENTREGABLE / "TFM_VISIBILIDAD_DETECCION_FP_WAZUH_FINAL_26062026.xlsx"
BEN_MASTER = ENTREGABLE / "TFM_BENCHMARK_RENDIMIENTO_VELOCIRAPTOR_FINAL_26062026.xlsx"
VIS_FINAL = ENTREGABLE / "TFM_VISIBILIDAD_DETECCION_FP_WAZUH_DEFINITIVO_20260712.xlsx"
BEN_FINAL = ENTREGABLE / "TFM_BENCHMARK_RENDIMIENTO_VELOCIRAPTOR_DEFINITIVO_20260712.xlsx"
OUTPUT = AUDIT_DIR / "AUDITORIA_CRUZADA_FINAL_20260712.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def check(results: list[dict[str, Any]], name: str, condition: bool, observed: Any) -> None:
    results.append({"check": name, "ok": bool(condition), "observed": observed})


def cell_map(worksheet: Any) -> dict[str, Any]:
    return {
        cell.coordinate: cell.value
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    }


def unchanged_sheets(master: Any, final: Any, names: list[str]) -> dict[str, Any]:
    details: dict[str, Any] = {}
    for name in names:
        before = cell_map(master[name])
        after = cell_map(final[name])
        changed = sorted(
            coordinate
            for coordinate in before.keys() | after.keys()
            if before.get(coordinate) != after.get(coordinate)
        )
        details[name] = {"unchanged": not changed, "changed_cells": changed[:25], "changed_count": len(changed)}
    return details


def series_ref(series: Any, attribute: str) -> str | None:
    value = getattr(series, attribute, None)
    if value is None:
        return None
    for ref_name in ("numRef", "strRef"):
        ref = getattr(value, ref_name, None)
        if ref is not None:
            return getattr(ref, "f", None)
    return None


def main() -> None:
    results: list[dict[str, Any]] = []
    vis = load_workbook(VIS_FINAL, data_only=False, read_only=False)
    ben = load_workbook(BEN_FINAL, data_only=False, read_only=False)
    vis_master = load_workbook(VIS_MASTER, data_only=False, read_only=False)
    ben_master = load_workbook(BEN_MASTER, data_only=False, read_only=False)
    try:
        check(results, "visibilidad_31_hojas", len(vis.sheetnames) == 31, len(vis.sheetnames))
        check(results, "benchmark_11_hojas", len(ben.sheetnames) == 11, len(ben.sheetnames))
        check(results, "wazuh_integrado", {"WAZUH_DETALLE", "COMPARACION_VR_WAZUH"}.issubset(vis.sheetnames), vis.sheetnames)
        check(results, "publicos_integrados", {"04_Control_Publicos", "15_Publicos_Definitivo"}.issubset(vis.sheetnames), vis.sheetnames)

        dashboard = vis["01_Dashboard"]
        check(results, "dashboard_total_tec_516", dashboard["B8"].value == 516, dashboard["B8"].value)
        check(results, "dashboard_total_fp_127", dashboard["B9"].value == 127, dashboard["B9"].value)
        check(results, "dashboard_hayabusa_3_9", dashboard["B12"].value == "3/9", dashboard["B12"].value)
        check(results, "dashboard_custom_9_9", dashboard["B13"].value == "9/9", dashboard["B13"].value)
        check(results, "dashboard_wazuh_4_9", dashboard["B14"].value == "4/9", dashboard["B14"].value)
        check(results, "dashboard_tec009_http_200", "HTTP 200" in str(dashboard["B15"].value), dashboard["B15"].value)

        public = vis["15_Publicos_Definitivo"]
        public_rows = list(public.iter_rows(min_row=5, max_row=10, values_only=True))
        public_by_name = {str(row[0]): row for row in public_rows}
        tec_sum = sum(int(row[12] or 0) for row in public_rows)
        fp_sum = sum(int(row[13] or 0) for row in public_rows)
        check(results, "publicos_seis_artifacts", len(public_by_name) == 6, sorted(public_by_name))
        check(results, "publicos_total_tec_516", tec_sum == 516, tec_sum)
        check(results, "publicos_total_fp_127", fp_sum == 127, fp_sum)
        check(results, "etw_no_concluyente", public_by_name["Windows.ETW.Monitoring"][17] == "NO CONCLUYENTE", public_by_name["Windows.ETW.Monitoring"][17])
        service = public_by_name["Windows.Events.ServiceCreation"]
        check(results, "service_creation_visibilidad_no_deteccion", str(service[4]).startswith("No:") and service[5] == "No", [service[4], service[5]])
        track = public_by_name["Generic.Events.TrackNetworkConnections"]
        check(results, "tracknetwork_128_12", track[12] == 128 and track[13] == 12, [track[12], track[13]])
        check(results, "tracknetwork_sin_atribucion", "PID=0" in str(track[18]) and "No atribuir" in str(track[18]), track[18])
        check(results, "tec009_exito_separado", "HTTP 200" in str(track[18]) and "UploadSucceeded=True" in str(track[18]), track[18])

        control_text = "\n".join(str(cell.value) for row in vis["04_Control_Publicos"].iter_rows() for cell in row if cell.value is not None)
        check(results, "chm_no_ejecutada", "NO EJECUTADA" in control_text and "SIN CAMPAÑA CHM" in control_text, "NO EJECUTADA / SIN CAMPAÑA CHM")
        hay = vis["13_Hayabusa_Resumen"]
        check(results, "hayabusa_185_matches", hay["B23"].value == 185, hay["B23"].value)
        check(results, "hayabusa_3_9_texto", hay["B25"].value == "3/9", hay["B25"].value)
        check(results, "hayabusa_fp_73", hay["B27"].value == 73, hay["B27"].value)
        check(results, "hayabusa_chm_no_existe", hay["B29"].value == "NO EXISTE", hay["B29"].value)

        executive = vis["RESUMEN_EJECUTIVO"]
        custom_values = {"total": executive["B15"].value, "P1": executive["B16"].value, "P2": executive["B17"].value, "P3": executive["B18"].value, "P4": executive["B19"].value}
        check(results, "custom_p1_p4_conservado", custom_values == {"total": 379, "P1": 19, "P2": 118, "P3": 109, "P4": 133}, custom_values)
        wazuh_values = {"base": executive["B21"].value, "custom": executive["B22"].value, "rules": executive["B23"].value, "gaps": executive["B24"].value}
        check(results, "wazuh_resultados_conservados", wazuh_values == {"base": 58, "custom": 110, "rules": 5, "gaps": 1}, wazuh_values)

        dashboard_chart_count = len(vis["01_Dashboard"]._charts)
        graph_charts = vis["GRAFICAS"]._charts
        total_vis_charts = sum(len(vis[name]._charts) for name in vis.sheetnames)
        check(results, "visibilidad_11_graficos", total_vis_charts == 11 and dashboard_chart_count == 6, {"total": total_vis_charts, "dashboard": dashboard_chart_count})
        wazuh_chart = graph_charts[4]
        wazuh_series = list(wazuh_chart.ser)
        wazuh_refs = {"categories": series_ref(wazuh_series[0], "cat"), "values": series_ref(wazuh_series[0], "val")} if wazuh_series else {}
        check(results, "wazuh_chart_una_serie", len(wazuh_series) == 1, len(wazuh_series))
        check(results, "wazuh_chart_rangos", wazuh_refs.get("categories") == "GRAFICAS!$A$35:$A$41" and wazuh_refs.get("values") == "GRAFICAS!$B$35:$B$41", wazuh_refs)

        vis_unchanged_names = [
            "02_Seleccion_Tecnicas", "07_Catalogo_Artifacts", "09_Benchmark_Plan", "10_Gaps_Custom_P1P4",
            "99_Listas", "11_Alertabilidad", "12_Plan_Alertas", "14_Arquitectura_Custom", "16_Scripts_Validados",
            "TECNICAS_REAL_VR", "VISIBILIDAD_SISTEMA", "ALERTAS_VR", "FP_RUNNER", "FP_HITS", "WAZUH_DETALLE"
        ]
        vis_unchanged = unchanged_sheets(vis_master, vis, vis_unchanged_names)
        check(results, "hojas_previas_criticas_conservadas", all(item["unchanged"] for item in vis_unchanged.values()), vis_unchanged)

        ben_exec = ben["RESUMEN_EJECUTIVO"]
        check(results, "benchmark_9_9_texto", ben_exec["C3"].value == "9/9", ben_exec["C3"].value)
        check(results, "benchmark_3_3_texto", ben_exec["B5"].value == "3/3", ben_exec["B5"].value)
        runs = ben["RUNS_VALIDOS"]
        headers = {str(cell.value): cell.column for cell in runs[1] if cell.value is not None}
        scenario_col = headers["Scenario"]
        repetition_col = headers["Repetition"]
        rows = [row for row in range(2, runs.max_row + 1) if runs.cell(row, scenario_col).value]
        scenarios = Counter(str(runs.cell(row, scenario_col).value) for row in rows)
        repetitions = Counter(str(runs.cell(row, repetition_col).value) for row in rows)
        check(results, "benchmark_nueve_runs", len(rows) == 9, len(rows))
        check(results, "benchmark_tres_escenarios_tres_runs", scenarios == Counter({"BASELINE_NO_VR": 3, "VR_IDLE": 3, "VR_TEC_RUNNER": 3}), scenarios)
        check(results, "benchmark_tres_repeticiones", repetitions == Counter({"1": 3, "2": 3, "3": 3}), repetitions)
        check(results, "server_gui_excluido", "SERVER_GUI" not in scenarios, scenarios)
        quality_text = "\n".join(str(cell.value) for row in ben["CALIDAD_DATOS"].iter_rows() for cell in row if cell.value is not None)
        check(results, "notepad_cero", "notepad.exe runner = 0" in quality_text, "notepad.exe runner = 0")
        ben_charts = ben["GRAFICAS"]._charts
        check(results, "benchmark_seis_graficos", len(ben_charts) == 6 and all(len(chart.ser) > 0 for chart in ben_charts), [len(chart.ser) for chart in ben_charts])
        check(results, "benchmark_grafico6_tres_series", len(ben_charts[5].ser) == 3, len(ben_charts[5].ser))

        ben_unchanged_names = ["RUNS_VALIDOS", "RESUMEN_ESCENARIO", "RAW_SAMPLES", "PROCESS_SAMPLES", "TRAZABILIDAD", "LIMITACIONES"]
        ben_unchanged = unchanged_sheets(ben_master, ben, ben_unchanged_names)
        check(results, "benchmark_metricas_fuente_conservadas", all(item["unchanged"] for item in ben_unchanged.values()), ben_unchanged)

        failures = [item for item in results if not item["ok"]]
        payload = {
            "generated_at": datetime.now().astimezone().isoformat(),
            "status": "APTO" if not failures else "NO APTO",
            "files": {
                "visibility": {"path": str(VIS_FINAL), "size": VIS_FINAL.stat().st_size, "sha256": sha256(VIS_FINAL)},
                "benchmark": {"path": str(BEN_FINAL), "size": BEN_FINAL.stat().st_size, "sha256": sha256(BEN_FINAL)},
            },
            "checks_total": len(results),
            "checks_ok": len(results) - len(failures),
            "checks_failed": len(failures),
            "checks": results,
        }
        OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
        print(json.dumps({"status": payload["status"], "checks_total": payload["checks_total"], "checks_failed": payload["checks_failed"], "output": str(OUTPUT)}, ensure_ascii=False))
        if failures:
            print(json.dumps(failures, ensure_ascii=False, indent=2, default=str))
            raise SystemExit(1)
    finally:
        vis.close()
        ben.close()
        vis_master.close()
        ben_master.close()


if __name__ == "__main__":
    main()
