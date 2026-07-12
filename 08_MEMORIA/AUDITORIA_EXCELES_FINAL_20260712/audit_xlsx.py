#!/usr/bin/env python3
"""Auditoría de solo lectura para libros XLSX del cierre del TFM."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


FORMULA_ERROR_TOKENS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A")
SHEET_REF_RE = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_. ]*))!")
CHART_FORMULA_RE = re.compile(r"<(?:\w+:)?f>(.*?)</(?:\w+:)?f>", re.DOTALL)
SERIES_RE = re.compile(r"<(?:\w+:)?ser(?:\s|>)")


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest().upper()


def safe_scalar(value: Any, max_len: int = 300) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    text = str(value)
    return text if len(text) <= max_len else text[: max_len - 1] + "…"


def chart_xml_audit(zf: zipfile.ZipFile) -> list[dict[str, Any]]:
    charts: list[dict[str, Any]] = []
    for name in sorted(n for n in zf.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", n)):
        raw = zf.read(name)
        text = raw.decode("utf-8", errors="replace")
        formulas = [re.sub(r"\s+", " ", f).strip() for f in CHART_FORMULA_RE.findall(text)]
        charts.append(
            {
                "part": name,
                "bytes": len(raw),
                "series_count": len(SERIES_RE.findall(text)),
                "formula_refs": formulas,
                "broken_refs": [f for f in formulas if "#REF!" in f.upper()],
                "has_title": bool(re.search(r"<(?:\w+:)?title(?:\s|>)", text)),
            }
        )
    return charts


def zip_audit(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {
        "is_zip": False,
        "zip_test": None,
        "entries": 0,
        "xml_parts": 0,
        "xml_parse_errors": [],
        "charts": [],
        "external_link_parts": [],
        "calc_chain_present": False,
        "vba_present": False,
    }
    try:
        with zipfile.ZipFile(path, "r") as zf:
            result["is_zip"] = True
            result["entries"] = len(zf.infolist())
            result["zip_test"] = zf.testzip()
            names = zf.namelist()
            result["xml_parts"] = sum(1 for n in names if n.endswith((".xml", ".rels")))
            result["external_link_parts"] = sorted(n for n in names if n.startswith("xl/externalLinks/"))
            result["calc_chain_present"] = "xl/calcChain.xml" in names
            result["vba_present"] = any(n.lower().endswith("vbaproject.bin") for n in names)
            for name in names:
                if not name.endswith((".xml", ".rels")):
                    continue
                try:
                    ET.fromstring(zf.read(name))
                except Exception as exc:  # pragma: no cover - defensive
                    result["xml_parse_errors"].append({"part": name, "error": repr(exc)})
            result["charts"] = chart_xml_audit(zf)
    except Exception as exc:
        result["fatal_error"] = repr(exc)
    return result


def table_record(table: Any) -> dict[str, Any]:
    return {
        "name": getattr(table, "name", None),
        "display_name": getattr(table, "displayName", None),
        "ref": getattr(table, "ref", None),
        "style": getattr(getattr(table, "tableStyleInfo", None), "name", None),
        "show_totals": bool(getattr(table, "totalsRowShown", False)),
    }


def workbook_audit(path: Path, previews: int = 18, preview_cols: int = 24) -> dict[str, Any]:
    stat = path.stat()
    report: dict[str, Any] = {
        "path": str(path.resolve()),
        "name": path.name,
        "size": stat.st_size,
        "modified": datetime.fromtimestamp(stat.st_mtime, timezone.utc).astimezone().isoformat(),
        "sha256": sha256_file(path),
        "openxml": zip_audit(path),
        "openpyxl": {},
    }

    formula_wb = load_workbook(path, read_only=False, data_only=False, keep_links=True)
    value_wb = load_workbook(path, read_only=False, data_only=True, keep_links=True)
    sheets: list[dict[str, Any]] = []
    total_formulas = 0
    total_nonempty = 0
    formula_token_errors: list[dict[str, str]] = []
    stored_errors: list[dict[str, str]] = []
    missing_sheet_refs: list[dict[str, str]] = []
    sheet_names = set(formula_wb.sheetnames)

    for ws in formula_wb.worksheets:
        data_ws = value_wb[ws.title]
        nonempty = 0
        formulas = 0
        cell_types: Counter[str] = Counter()
        previews_out: list[dict[str, Any]] = []
        for row in ws.iter_rows():
            row_preview: list[dict[str, Any]] = []
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                nonempty += 1
                cell_types[str(cell.data_type)] += 1
                if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                    formulas += 1
                    text = str(value)
                    for token in FORMULA_ERROR_TOKENS:
                        if token.upper() in text.upper():
                            formula_token_errors.append({"sheet": ws.title, "cell": cell.coordinate, "formula": text})
                            break
                    for match in SHEET_REF_RE.finditer(text):
                        ref_sheet = (match.group(1) or match.group(2) or "").replace("''", "'").strip()
                        if ref_sheet and ref_sheet not in sheet_names:
                            missing_sheet_refs.append(
                                {"sheet": ws.title, "cell": cell.coordinate, "referenced_sheet": ref_sheet, "formula": text}
                            )
                if len(previews_out) < previews and len(row_preview) < preview_cols:
                    row_preview.append({"cell": cell.coordinate, "value": safe_scalar(value)})
            if row_preview and len(previews_out) < previews:
                previews_out.append({"row": row[0].row if row else None, "cells": row_preview})

        for row in data_ws.iter_rows():
            for cell in row:
                if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value in FORMULA_ERROR_TOKENS):
                    stored_errors.append({"sheet": ws.title, "cell": cell.coordinate, "value": str(cell.value)})

        tables = [table_record(t) for t in ws.tables.values()]
        charts = []
        for chart in ws._charts:
            charts.append(
                {
                    "type": type(chart).__name__,
                    "title_present": getattr(chart, "title", None) is not None,
                    "series_count": len(getattr(chart, "ser", []) or []),
                    "anchor": safe_scalar(getattr(chart, "anchor", None)),
                }
            )
        sheet = {
            "name": ws.title,
            "state": ws.sheet_state,
            "dimension": ws.calculate_dimension(),
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "max_column_letter": get_column_letter(ws.max_column),
            "nonempty_cells": nonempty,
            "formula_cells": formulas,
            "cell_types": dict(cell_types),
            "tables": tables,
            "charts": charts,
            "images": len(ws._images),
            "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "auto_filter": ws.auto_filter.ref,
            "print_area": str(ws.print_area) if ws.print_area else None,
            "data_validations": len(ws.data_validations.dataValidation),
            "preview": previews_out,
        }
        sheets.append(sheet)
        total_formulas += formulas
        total_nonempty += nonempty

    defined_names = []
    for item in formula_wb.defined_names.values():
        defined_names.append(
            {
                "name": getattr(item, "name", None),
                "attr_text": getattr(item, "attr_text", None),
                "local_sheet_id": getattr(item, "localSheetId", None),
                "hidden": bool(getattr(item, "hidden", False)),
            }
        )

    calc = formula_wb.calculation
    report["openpyxl"] = {
        "sheet_count": len(formula_wb.sheetnames),
        "sheet_names": formula_wb.sheetnames,
        "sheets": sheets,
        "total_nonempty_cells": total_nonempty,
        "total_formula_cells": total_formulas,
        "total_tables": sum(len(s["tables"]) for s in sheets),
        "total_charts": sum(len(s["charts"]) for s in sheets),
        "total_images": sum(s["images"] for s in sheets),
        "defined_names": defined_names,
        "external_links": len(getattr(formula_wb, "_external_links", []) or []),
        "formula_token_errors": formula_token_errors,
        "stored_error_cells": stored_errors,
        "missing_sheet_references": missing_sheet_refs,
        "calculation": {
            "calc_mode": getattr(calc, "calcMode", None),
            "full_calc_on_load": getattr(calc, "fullCalcOnLoad", None),
            "force_full_calc": getattr(calc, "forceFullCalc", None),
            "calc_on_save": getattr(calc, "calcOnSave", None),
            "calc_id": getattr(calc, "calcId", None),
        },
        "properties": {
            "title": formula_wb.properties.title,
            "subject": formula_wb.properties.subject,
            "creator": formula_wb.properties.creator,
            "last_modified_by": formula_wb.properties.lastModifiedBy,
            "created": safe_scalar(formula_wb.properties.created),
            "modified": safe_scalar(formula_wb.properties.modified),
            "keywords": formula_wb.properties.keywords,
        },
    }
    formula_wb.close()
    value_wb.close()
    return report


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbooks", nargs="+", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--previews", type=int, default=18)
    parser.add_argument("--preview-cols", type=int, default=24)
    parser.add_argument("--quiet", action="store_true")
    args = parser.parse_args()
    reports = [workbook_audit(p, args.previews, args.preview_cols) for p in args.workbooks]
    payload = {
        "generated_at": datetime.now(timezone.utc).astimezone().isoformat(),
        "python": sys.version,
        "openpyxl": __import__("openpyxl").__version__,
        "workbooks": reports,
    }
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text + os.linesep, encoding="utf-8")
    if not args.quiet:
        print(text)
    else:
        print(
            json.dumps(
                [
                    {
                        "path": r["path"],
                        "size": r["size"],
                        "sha256": r["sha256"],
                        "sheets": r["openpyxl"]["sheet_count"],
                        "tables": r["openpyxl"]["total_tables"],
                        "charts": r["openpyxl"]["total_charts"],
                        "formulas": r["openpyxl"]["total_formula_cells"],
                    }
                    for r in reports
                ],
                ensure_ascii=False,
            )
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
